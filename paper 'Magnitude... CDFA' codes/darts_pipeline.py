"""The Darts pipeline: KNN / LSTM with and without CDFA, 8:2 direct forecast.

Raw-unit MAE, no scaling.  The table lists the four models computed here
first, followed by published baselines for the same datasets, which are shown
for context only and take no part in the computation.
"""
import os

import numpy as np
import pandas as pd

import cdfa
from cdfa import ALPHA, corr_cdfa_weights
from datasets import DART_NAMES, SPLIT, load_darts
from models import knn_direct, run_lstm, tune_knn

OURS = ["KNN", "CDFA+KNN", "LSTM", "CDFA+LSTM"]

# Published baselines for the same datasets; display only.
PAPER = {
    "AirPassengersDataset": [34.67, 24.03, 54.96, 97.89, 59.16, 34.37, 62.51, 44.65, 81.45],
    "AusBeerDataset":       [102.05, 17.13, 30.90, 10.39, 34.23, 16.13, 11.94, 21.97, 96.35],
    "GasRateCO2Dataset":    [2.27, 2.37, 2.64, 2.63, 3.85, 3.50, 2.50, 2.67, 2.29],
    "MonthlyMilkDataset":   [30.33, 37.19, 70.86, 33.64, 32.73, 9.68, 28.09, 42.60, 85.71],
    "SunspotsDataset":      [53.74, 43.56, 51.82, 73.15, 49.93, 47.34, 41.40, 62.33, 48.24],
    "WineDataset":          [4552.06, 2306.70, 3287.14, 4562.02, 3909.51, 1569.32, 2871.33, 2498.69, 4075.28],
    "WoolyDataset":         [649.98, 588.78, 1158.79, 903.01, 382.09, 808.73, 728.92, 542.28, 1210.33],
    "HeartRateDataset":     [5.65, 5.56, 5.49, 6.57, 6.10, 6.21, 5.85, 6.74, 5.92],
}
PAPER_COLS = ["GP", "ARIMA", "TCN", "N-BEATS", "N-HiTS", "llmtime(ZS)",
              "TimesFM(ZS)", "PatchTST", "NAIVE"]
PAPER_AM = [0.819, 0.6045, 0.8426, 0.9178, 0.8106, 0.6643, 0.683, 0.7414, 1.0]
PAPER_GM = [0.7507, 0.5219, 0.7944, 0.7316, 0.6935, 0.4883, 0.5768, 0.6431, 1.0]


def collect():
    # Run the 8:2 split on every Darts dataset and collect the raw MAE.
    rows = []
    for name in DART_NAMES:
        y = load_darts(name)
        k = int(len(y) * SPLIT)
        train, test = y[:k], y[k:]
        r = {"Dataset": name}
        for tag, flag in [("KNN", False), ("CDFA+KNN", True)]:
            L, kk = tune_knn(train, len(test), flag)
            if L is None:
                L, kk = 2, 1
            aug = (corr_cdfa_weights(train, L), ALPHA) if flag else None
            r[tag] = float(np.abs(knn_direct(train, len(test), L, kk, aug)
                                  - test).mean())
        for tag, flag in [("LSTM", False), ("CDFA+LSTM", True)]:
            r[tag] = run_lstm(train, test, flag)
        r["NAIVE"] = float(np.abs(test - train[-1]).mean())
        rows.append(r)
        print(f"{name:<22} KNN={r['KNN']:.4f} CDFA+KNN={r['CDFA+KNN']:.4f} "
              f"LSTM={r['LSTM']:.4f} CDFA+LSTM={r['CDFA+LSTM']:.4f} "
              f"NAIVE={r['NAIVE']:.4f}")
    return pd.DataFrame(rows)


def scaled_rows(df):
    # Mean and geometric mean of MAE/NAIVE across datasets.
    sc = {c: (df[c] / df["NAIVE"]) for c in OURS}
    am = [float(np.mean(sc[c])) for c in OURS]
    gm = [float(np.exp(np.mean(np.log(sc[c])))) for c in OURS]
    return am, gm


def print_table(df, am, gm):
    w = 36
    PW = 13
    head = "Dataset".ljust(w) + "".join(c.rjust(12) for c in OURS) + "NAIVE".rjust(12)
    head += "  ||" + "".join(c.rjust(PW) for c in PAPER_COLS)
    print("\n" + "=" * len(head))
    print("MAE  (8:2 direct / only train once ; no TW)".center(len(head)))
    print("=" * len(head))
    print(head)
    print("-" * len(head))
    for _, r in df.iterrows():
        line = str(r["Dataset"]).ljust(w)
        line += "".join(f"{r[c]:12.4f}" for c in OURS) + f"{r['NAIVE']:12.4f}"
        line += "  ||" + "".join(f"{v:>{PW}.2f}" for v in PAPER[r["Dataset"]])
        print(line)
    print("-" * len(head))
    for label, vals in [("Scaled MAE (Arithmetic Mean)", am),
                        ("Scaled MAE (Geometric Mean)", gm)]:
        line = label.ljust(w) + "".join(f"{v:12.4f}" for v in vals) + f"{1.0:12.4f}"
        line += "  ||" + "".join(f"{v:>{PW}.4f}" for v in
                                 (PAPER_AM if "Arith" in label else PAPER_GM))
        print(line)
    print("-" * len(head))
    verdict = ["√" if am[1] < am[0] else "×", "", "√" if am[3] < am[2] else "×", ""]
    print("is CDFA better than without CDFA?".ljust(w)
          + "".join(v.rjust(12) for v in verdict))
    print("=" * len(head))


def save_xlsx(df, am, gm, path):
    try:
        import openpyxl
    except ImportError:
        import subprocess, sys
        subprocess.check_call([sys.executable, "-m", "pip", "install",
                               "--quiet", "openpyxl>=3.1.0"])
        import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    HEAD = "FF4472C4"
    thin = Side(style="thin", color="FFBFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    ws = wb.active
    ws.title = "MAE_8-2"
    cols = ["Dataset"] + OURS + ["NAIVE"] + PAPER_COLS
    for c, name in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=HEAD)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = box
    for i, (_, r) in enumerate(df.iterrows()):
        row = 2 + i
        ws.cell(row=row, column=1, value=r["Dataset"])
        for c, col in enumerate(OURS + ["NAIVE"], 2):
            ws.cell(row=row, column=c, value=float(r[col]))
        for c, v in enumerate(PAPER[r["Dataset"]], 2 + len(OURS) + 1):
            ws.cell(row=row, column=c, value=float(v))
    last = 1 + len(df)
    for c, name in enumerate(["Scaled MAE (Arithmetic Mean)",
                              "Scaled MAE (Geometric Mean)"], 1):
        row = last + c
        ws.cell(row=row, column=1, value=name)
        vals = am if c == 1 else gm
        for i2, v in enumerate(vals, 2):
            ws.cell(row=row, column=i2, value=v)
        ws.cell(row=row, column=2 + len(OURS), value=1.0)
        for i2, v in enumerate(PAPER_AM if c == 1 else PAPER_GM,
                               2 + len(OURS) + 1):
            ws.cell(row=row, column=i2, value=v)
    vrow = last + 3
    for row in range(1, vrow + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=row, column=c)
            cell.border = box
            if c > 1 and isinstance(cell.value, float):
                cell.number_format = "#,##0.0000"
                cell.alignment = Alignment(horizontal="center")
    for c in range(1, len(cols) + 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 15
    ws.column_dimensions["A"].width = 30
    ws.freeze_panes = "B2"
    wb.properties.title = "knn_lstm_cdfa_mae_table"
    wb.save(path)


def run(out_dir='outputs'):
    # Entry point used by main.py.
    os.makedirs(out_dir, exist_ok=True)
    df = collect()
    am, gm = scaled_rows(df)
    print_table(df, am, gm)

    out = df.copy()
    for c in OURS + ["NAIVE"]:
        out[c + "_scaled"] = out[c] / out["NAIVE"]
    out.to_csv(os.path.join(out_dir, "knn_lstm_cdfa_mae_table.csv"),
               index=False)
    try:
        save_xlsx(df, am, gm, os.path.join(out_dir,
                                           "knn_lstm_cdfa_mae_table.xlsx"))
        print("\nsaved: knn_lstm_cdfa_mae_table.csv / .xlsx")
    except Exception as e:
        print("\nxlsx skipped:", e)
    return df
