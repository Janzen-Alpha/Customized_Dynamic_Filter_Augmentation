import os
import numpy as np
import pandas as pd

# 输出原始量纲 MAE（不再除以 NAIVE），版式对齐 "TimesFM Darts result compare.xlsx"，
# 只保留 "only train once"（直接预测）那一组，不做 TW。
# 计算管线与 knn_lstm_cdfa_darts.py 完全一致。

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = r"D:\MPU\Program\Amang Research Lab\26.9.25 尝试CDFA在Darts数据集上\datasets\darts_csv"

SPLIT = 0.8
ALPHA = 0.3
L_GRID = [2, 3, 4, 6, 8, 12, 16, 24, 36, 48]
K_GRID = [1, 2, 3, 5, 8, 10]
LSTM_L = [12, 24]
HIDDEN, EPOCHS, LR, BATCH = 16, 300, 0.02, 64
SEEDS = [0, 1, 2, 3, 4]

NAMES = ["AirPassengersDataset", "AusBeerDataset", "GasRateCO2Dataset",
         "MonthlyMilkDataset", "SunspotsDataset", "WineDataset",
         "WoolyDataset", "HeartRateDataset"]

# 论文原值（参照列，取自 TimesFM Darts result compare.xlsx，不参与计算）
PAPER = {
    "AirPassengersDataset": [34.67, 24.03, 54.96, 97.89, 59.16, 34.37, 62.51, 44.65, 81.45],
    "AusBeerDataset":       [102.05, 17.13, 30.90, 10.39, 34.23, 16.13, 11.94, 21.97, 96.35],
    "GasRateCO2Dataset":    [2.27, 2.37, 2.64, 2.63, 3.85, 3.50, 2.50, 2.67, 2.29],
    "MonthlyMilkDataset":   [30.33, 37.19, 70.86, 33.64, 32.73, 9.68, 28.09, 42.60, 85.71],
    "SunspotsDataset":      [53.74, 43.56, 51.82, 73.15, 49.93, 47.34, 41.40, 62.33, 48.24],
    "WineDataset":          [4552.06, 2306.70, 3287.14, 4562.02, 3909.51, 1569.32, 2871.33, 2498.69, 4075.28],
    "WoolyDataset":         [649.98, 588.78, 1158.79, 903.01, 382.09, 808.73, 728.92, 542.28, 1210.33],
    "HeartRateDataset":     [5.65, 5.56, 5.49, 6.57, 6.10, 6.21, 5.85, 6.74, 5.92],
}
PAPER_COLS = ["GP", "ARIMA", "TCN", "N-BEATS", "N-HiTS", "llmtime(ZS)", "TimesFM(ZS)", "PatchTST", "NAIVE"]
PAPER_AM = [0.819, 0.6045, 0.8426, 0.9178, 0.8106, 0.6643, 0.683, 0.7414, 1.0]
PAPER_GM = [0.7507, 0.5219, 0.7944, 0.7316, 0.6935, 0.4883, 0.5768, 0.6431, 1.0]

OURS = ["KNN", "CDFA+KNN", "LSTM", "CDFA+LSTM"]


def load(name):
    """读一个 Darts 数据集，返回 1 维序列；Sunspots 取每 4 期、HeartRate 取每 2 期。"""
    df = pd.read_csv(os.path.join(DATA_DIR, name + ".csv"))
    col = 2 if name == "GasRateCO2Dataset" else 1
    s = df.iloc[:, col].astype(float).reset_index(drop=True)
    if name == "SunspotsDataset":
        s = s.iloc[::4].reset_index(drop=True)
    if name == "HeartRateDataset":
        s = s.iloc[::2].reset_index(drop=True)
    return s.to_numpy()


def minmax(s):
    """线性缩放到 [0,1]；常数列返回全 0。"""
    s = np.asarray(s, float)
    lo, hi = s.min(), s.max()
    return (s - lo) / (hi - lo) if hi - lo > 1e-12 else np.zeros_like(s)


def xi_returns(s, w, k):
    """论文 Eq.7 局部相关场：在一阶差分上算 lag=1..w 的自相关，返回长度 w 向量（只用 s[:k+1]）。"""
    r = np.diff(np.asarray(s[:k + 1], float))
    out = np.zeros(w)
    if len(r) < w + 2:
        return out
    for lag in range(1, w + 1):
        if len(r) <= lag + 1:
            continue
        x, y = r[lag:], r[:-lag]
        xm, ym = x.mean(), y.mean()
        num = ((x - xm) * (y - ym)).sum()
        den = np.sqrt(((x - xm) ** 2).sum() * ((y - ym) ** 2).sum())
        out[lag - 1] = num / den if den > 1e-12 else 0.0
    return out


def phi_signed(xi):
    """论文 Eq.13 分段 min-max：按符号分两支映射到 [-1,1]。"""
    hi, lo = xi.max(), xi.min()
    d = hi - lo
    if d < 1e-12:
        return np.zeros(len(xi))
    return np.array([(v - lo) / d if v >= 0 else (v - hi) / d for v in xi])


def cdfa_weights(hist, w):
    """从 hist 算 CDFA 权重；hist 不足 w+1 点时退化为均匀核 1/w。"""
    if len(hist) - 1 < w:
        return np.ones(w) / w
    return phi_signed(xi_returns(minmax(hist), w, len(hist) - 1))


def perturb(X, phi, alpha):
    """CDFA 差分对搬移：只用于库窗口，不碰 target / 查询窗。"""
    X = np.asarray(X, float)
    if X.ndim == 1:
        X = X[None, :]
    m = min(len(phi), X.shape[1] - 1)
    if m <= 0:
        return X.copy()
    d = np.diff(X, axis=1)
    a = np.zeros_like(X)
    a[:, :m] = d[:, :m] * np.asarray(phi[:m]) * alpha
    adj = 0.5 * a
    adj[:, 1:] += 0.5 * a[:, :-1]
    return X + adj


def knn_direct(hist, H, L, k, aug=None):
    """直接多步：库窗口 hist[t-L:t] -> 其后 H 点；取最近 k 个邻居的 H 点均值。aug=(phi,alpha) 时先增强库。"""
    n = len(hist)
    X = [hist[t - L:t] for t in range(L, n - H + 1)]
    T = [hist[t:t + H] for t in range(L, n - H + 1)]
    if not X:
        return np.full(H, hist[-1])
    X, T = np.array(X, float), np.array(T, float)
    if aug is not None:
        X = perturb(X, aug[0], aug[1])
    q = np.asarray(hist, float)[-L:]
    d = np.sqrt(((X - q) ** 2).sum(axis=1))
    return T[np.argsort(d)[:k]].mean(axis=0)


def tune_knn(train, H, use_cdfa):
    """只用 train 调 (L, k)：验证窗取 train 末尾 h 点，回测 MAE 最小的一组。"""
    n = len(train)
    best = (np.inf, None, None)
    for L in L_GRID:
        if L > n - H:
            continue
        h = min(min(H, max(2, n - 10)), max(2, (n - L) // 2))
        if h < 2:
            continue
        hist, true = train[:-h], train[-h:]
        if L > len(hist) - 3 or len(hist) - L - h + 1 < 1:
            continue
        aug = (cdfa_weights(hist, L), ALPHA) if use_cdfa else None
        for k in K_GRID:
            if k > len(hist) - L:
                continue
            mae = np.abs(knn_direct(hist, h, L, k, aug) - true).mean()
            if mae < best[0] - 1e-12:
                best = (mae, L, k)
    return best[1], best[2]


class LSTM:
    """NumPy 手写单层 vanilla LSTM（前向 + BPTT + Adam）；d_in/d_h/d_out 为输入/隐层/输出维。"""

    def __init__(self, d_in, d_h, d_out, seed=0):
        """初始化四门权重与输出层；seed 固定随机性，并建好 Adam 动量缓存。"""
        rng = np.random.RandomState(seed)
        self.d_in, self.d_h, self.d_out = d_in, d_h, d_out
        k = 1.0 / np.sqrt(d_in + d_h)
        self.W = rng.uniform(-k, k, (d_in + d_h, 4 * d_h))
        self.b = np.zeros(4 * d_h)
        self.b[d_h:2 * d_h] = 1.0
        self.Wy = rng.uniform(-k, k, (d_h, d_out))
        self.by = np.zeros(d_out)
        self.m = {n: np.zeros_like(getattr(self, n)) for n in ('W', 'b', 'Wy', 'by')}
        self.v = {n: np.zeros_like(getattr(self, n)) for n in ('W', 'b', 'Wy', 'by')}
        self.step = 0

    def forward(self, X, cache=True):
        """X 形状 (batch, 时间步, d_in)；返回最后隐状态经 Wy 映射的 d_out 维输出。"""
        B, T, _ = X.shape
        hd = self.d_h
        h = np.zeros((B, hd)); c = np.zeros((B, hd)); store = []
        for t in range(T):
            z = np.concatenate([X[:, t, :], h], axis=1)
            q = z @ self.W + self.b
            i = 1 / (1 + np.exp(-np.clip(q[:, 0:hd], -30, 30)))
            f = 1 / (1 + np.exp(-np.clip(q[:, hd:2 * hd], -30, 30)))
            g = np.tanh(q[:, 2 * hd:3 * hd])
            o = 1 / (1 + np.exp(-np.clip(q[:, 3 * hd:4 * hd], -30, 30)))
            cp = c
            c = f * cp + i * g
            h = o * np.tanh(c)
            if cache:
                store.append((z, i, f, g, o, c, cp))
        self.hT, self.cache = h, store
        return h @ self.Wy + self.by

    def backward(self, X, dY):
        """BPTT：返回 (dW, db, dWy, dby)，需要 forward 留下的 cache。"""
        T = X.shape[1]; hd = self.d_h
        dW = np.zeros_like(self.W); db = np.zeros_like(self.b)
        dWy = self.hT.T @ dY; dby = dY.sum(0)
        dh = dY @ self.Wy.T
        dc = np.zeros((len(X), hd))
        for t in reversed(range(T)):
            z, i, f, g, o, c, cp = self.cache[t]
            th = np.tanh(c)
            do = dh * th
            dcc = dh * o * (1 - th * th) + dc
            di, dg, df = dcc * g, dcc * i, dcc * cp
            dc = dcc * f
            dq = np.concatenate([di * i * (1 - i), df * f * (1 - f),
                                 dg * (1 - g * g), do * o * (1 - o)], axis=1)
            dW += z.T @ dq; db += dq.sum(0)
            dh = (dq @ self.W.T)[:, self.d_in:]
        return dW, db, dWy, dby

    def adam(self, grads, lr, b1=0.9, b2=0.999, eps=1e-8, clip=5.0):
        """Adam 单步更新；lr 学习率，clip 梯度范数上限，eps 防除零。"""
        self.step += 1
        for n, gr in grads.items():
            nn = np.linalg.norm(gr)
            if nn > clip:
                gr = gr * (clip / nn)
            self.m[n] = b1 * self.m[n] + (1 - b1) * gr
            self.v[n] = b2 * self.v[n] + (1 - b2) * gr * gr
            mh = self.m[n] / (1 - b1 ** self.step)
            vh = self.v[n] / (1 - b2 ** self.step)
            setattr(self, n, getattr(self, n) - lr * mh / (np.sqrt(vh) + eps))

    def fit(self, X, Y, epochs, lr, batch, seed):
        """训练若干轮：打乱 + 分批复 forward/backward/Adam；seed 固定 shuffle。"""
        rng = np.random.RandomState(seed)
        for _ in range(epochs):
            idx = rng.permutation(len(X))
            for s in range(0, len(X), batch):
                b = idx[s:s + batch]
                dY = 2.0 * (self.forward(X[b]) - Y[b]) / len(b)
                self.adam(dict(zip(('W', 'b', 'Wy', 'by'), self.backward(X[b], dY))), lr)
        return self

    def predict(self, X):
        """推理（不缓存），返回 d_out 维预测。"""
        return self.forward(X, cache=False)


def lstm_mae(hist, true, L, H, use_cdfa, seed):
    """在 hist 上训练、对 true（长度 H）算原始 MAE；use_cdfa 时对训练窗口做增强。"""
    aug = (cdfa_weights(hist, L), ALPHA) if use_cdfa else None
    lo, hi = float(hist.min()), float(hist.max())
    if hi <= lo:
        hi = lo + 1.0
    sn = (np.asarray(hist, float) - lo) / (hi - lo)
    X = [sn[t - L:t] for t in range(L, len(sn) - H + 1)]
    Y = [sn[t:t + H] for t in range(L, len(sn) - H + 1)]
    if len(X) < 5:
        return np.inf
    X, Y = np.array(X, float), np.array(Y, float)
    if aug is not None:
        X = perturb(X, aug[0], aug[1])
    net = LSTM(1, HIDDEN, Y.shape[1], seed).fit(X[:, :, None], Y, EPOCHS, LR, BATCH, seed)
    p = net.predict(sn[-L:][None, :, None])[0]
    return float(np.abs(p * (hi - lo) + lo - true).mean())


def run_lstm(train, test, use_cdfa):
    """只用 train 挑窗长 L，再用 SEEDS 多个 seed 在 test 上取平均原始 MAE。"""
    n, H = len(train), len(test)
    best = (np.inf, LSTM_L[0])
    for L in LSTM_L:
        if L + H > n or n - L - H + 1 < 5:
            continue
        h = min(min(H, max(2, n - 40)), max(2, (n - L + 1 - 5) // 2))
        if h < 2 or n - L - h + 1 < 5:
            continue
        m = lstm_mae(train[:-h], train[-h:], L, h, use_cdfa, SEEDS[0])
        if np.isfinite(m) and m < best[0] - 1e-12:
            best = (m, L)
    L = best[1]
    if L + H > n or n - L - H + 1 < 5:
        return np.nan
    ms = np.array([lstm_mae(train, test, L, H, use_cdfa, s) for s in SEEDS])
    ms = ms[np.isfinite(ms)]
    return float(ms.mean()) if len(ms) else np.nan


def collect():
    """跑一遍 8:2，收集四个模型的原始量纲 MAE 与 NAIVE（都是 MAE，不除以 NAIVE）。"""
    rows = []
    for name in NAMES:
        y = load(name)
        k = int(len(y) * SPLIT)
        train, test = y[:k], y[k:]
        r = {"Dataset": name}
        for tag, flag in [("KNN", False), ("CDFA+KNN", True)]:
            L, kk = tune_knn(train, len(test), flag)
            if L is None:
                L, kk = 2, 1
            aug = (cdfa_weights(train, L), ALPHA) if flag else None
            r[tag] = float(np.abs(knn_direct(train, len(test), L, kk, aug) - test).mean())
        for tag, flag in [("LSTM", False), ("CDFA+LSTM", True)]:
            r[tag] = run_lstm(train, test, flag)
        r["NAIVE"] = float(np.abs(test - train[-1]).mean())
        rows.append(r)
        print(f"{name:<22} KNN={r['KNN']:.4f} CDFA+KNN={r['CDFA+KNN']:.4f} "
              f"LSTM={r['LSTM']:.4f} CDFA+LSTM={r['CDFA+LSTM']:.4f} NAIVE={r['NAIVE']:.4f}")
    return pd.DataFrame(rows)


def scaled_rows(df):
    """按 (MAE / NAIVE) 逐数据集缩放后取算术 / 几何平均，与参考表口径一致。"""
    sc = {c: (df[c] / df["NAIVE"]) for c in OURS}
    am = [float(np.mean(sc[c])) for c in OURS]
    gm = [float(np.exp(np.mean(np.log(sc[c])))) for c in OURS]
    return am, gm


def print_table(df, am, gm):
    """按参考表版式打印：数据集行 + Scaled MAE(AM/GM) + 判定行。"""
    w = 36   # 需容纳 "Scaled MAE (Arithmetic Mean)" 这类标签，保证数字列对齐
    PW = 13
    head = "Dataset".ljust(w) + "".join(c.rjust(12) for c in OURS) + "NAIVE".rjust(12)
    head += "  ||" + "".join(c.rjust(PW) for c in PAPER_COLS)
    print("\n" + "=" * len(head))
    print("MAE  (8:2 direct / only train once ; no TW)".center(len(head)))
    print("=" * len(head))
    print(head)
    print("-" * len(head))
    for _, r in df.iterrows():
        line = str(r["Dataset"]).ljust(w)
        line += "".join(f"{r[c]:12.4f}" for c in OURS) + f"{r['NAIVE']:12.4f}"
        line += "  ||" + "".join(f"{v:>{PW}.2f}" for v in PAPER[r["Dataset"]])
        print(line)
    print("-" * len(head))
    for label, vals in [("Scaled MAE (Arithmetic Mean)", am), ("Scaled MAE (Geometric Mean)", gm)]:
        line = label.ljust(w) + "".join(f"{v:12.4f}" for v in vals) + f"{1.0:12.4f}"
        line += "  ||" + "".join(f"{v:>{PW}.4f}" for v in (PAPER_AM if "Arith" in label else PAPER_GM))
        print(line)
    print("-" * len(head))
    verdict = ["√" if am[1] < am[0] else "×", "", "√" if am[3] < am[2] else "×", ""]
    print("is CDFA better than without CDFA?".ljust(w) + "".join(v.rjust(12) for v in verdict))
    print("=" * len(head))


def save_xlsx(df, am, gm, path):
    """把同一张表落成 xlsx（静态值，无公式）。"""
    try:
        import openpyxl
    except ImportError:
        import subprocess, sys
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
        import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    HEAD = "FF4472C4"
    thin = Side(style="thin", color="FFBFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    ws = wb.active
    ws.title = "MAE_8-2"
    cols = ["Dataset"] + OURS + ["NAIVE"] + PAPER_COLS
    for c, name in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=HEAD)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = box
    for i, (_, r) in enumerate(df.iterrows()):
        row = 2 + i
        ws.cell(row=row, column=1, value=r["Dataset"])
        for c, col in enumerate(OURS + ["NAIVE"], 2):
            ws.cell(row=row, column=c, value=float(r[col]))
        for c, v in enumerate(PAPER[r["Dataset"]], 2 + len(OURS) + 1):
            ws.cell(row=row, column=c, value=float(v))
    last = 1 + len(df)
    for c, name in enumerate(["Scaled MAE (Arithmetic Mean)", "Scaled MAE (Geometric Mean)"], 1):
        row = last + c
        ws.cell(row=row, column=1, value=name)
        vals = am if c == 1 else gm
        for i2, v in enumerate(vals, 2):
            ws.cell(row=row, column=i2, value=v)
        ws.cell(row=row, column=2 + len(OURS), value=1.0)
        for i2, v in enumerate(PAPER_AM if c == 1 else PAPER_GM, 2 + len(OURS) + 1):
            ws.cell(row=row, column=i2, value=v)
    vrow = last + 3
    ws.cell(row=vrow, column=1, value="is CDFA better than without CDFA?")
    ws.cell(row=vrow, column=2, value="√" if am[1] < am[0] else "×")
    ws.cell(row=vrow, column=4, value="√" if am[3] < am[2] else "×")
    for row in range(1, vrow + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=row, column=c)
            cell.border = box
            if c > 1 and isinstance(cell.value, float):
                cell.number_format = "#,##0.0000"
                cell.alignment = Alignment(horizontal="center")
    for c in range(1, len(cols) + 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 15
    ws.column_dimensions["A"].width = 30
    ws.freeze_panes = "B2"
    wb.properties.title = "knn_lstm_cdfa"
    wb.save(path)


if __name__ == "__main__":
    df = collect()
    am, gm = scaled_rows(df)
    print_table(df, am, gm)

    df_out = df.copy()
    for c in OURS + ["NAIVE"]:
        df_out[c + "_scaled"] = df_out[c] / df_out["NAIVE"]
    df_out.to_csv(os.path.join(HERE, "knn_lstm_cdfa.csv"), index=False)

    # try:
    #     save_xlsx(df, am, gm, os.path.join(HERE, "knn_lstm_cdfa.xlsx"))
    #     print("\nsaved: knn_lstm_cdfa.csv / knn_lstm_cdfa.xlsx")
    # except Exception as e:
    #     print("\nxlsx skipped:", e)
